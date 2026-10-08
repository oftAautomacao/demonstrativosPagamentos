import argparse
import datetime as dt
import hashlib
import json
import os
import sys
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from tuss_processor import (
    find_mapping_file,
    is_current_header,
    is_mapping_amount_header,
    is_new_header,
    normalize_text,
)


MAX_ROWS = 20_000
MAX_COLUMNS = 64


def is_within(path, root):
    try:
        return os.path.commonpath([str(path), str(root)]) == str(root)
    except ValueError:
        return False


def workbook_options(path):
    return {
        "read_only": False,
        "data_only": False,
        "keep_vba": path.suffix.lower() == ".xlsm",
    }


def open_workbook(path):
    stream = path.open("rb") if path.suffix.lower() == ".xslx" else None
    try:
        workbook = openpyxl.load_workbook(stream or path, **workbook_options(path))
    except Exception:
        if stream is not None:
            stream.close()
        raise
    return workbook, stream


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def table_identifier(root, path, sheet_name, header_row):
    relative = path.relative_to(root)
    source = f"{relative.as_posix()}|{sheet_name}|{header_row}"
    return hashlib.sha256(source.encode("utf-8")).hexdigest()[:16]


def matching_excel_table(worksheet, header_row, current_column, new_column):
    for table in worksheet.tables.values():
        min_column, min_row, max_column, max_row = range_boundaries(table.ref)
        if min_row == header_row and min_column <= current_column <= max_column and min_column <= new_column <= max_column:
            return table, min_column, max_column, max_row
    return None, None, None, None


def discover_tables(workbook, root, plan_dir, mapping_path):
    discovered = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            current_columns = [cell.column for cell in row if is_current_header(cell.value)]
            new_columns = [cell.column for cell in row if is_new_header(cell.value)]
            if not current_columns or not new_columns:
                continue

            header_row = row[0].row
            current_column = current_columns[0]
            new_column = new_columns[0]
            excel_table, min_column, max_column, end_row = matching_excel_table(
                worksheet, header_row, current_column, new_column
            )
            if excel_table is None:
                populated_columns = [cell.column for cell in row if cell.value not in (None, "")]
                min_column = min(populated_columns)
                max_column = max(populated_columns)
                end_row = header_row
                for row_index in range(header_row + 1, worksheet.max_row + 1):
                    if any(
                        worksheet.cell(row_index, column).value not in (None, "")
                        for column in range(min_column, max_column + 1)
                    ):
                        end_row = row_index

            columns = []
            for column in range(min_column, max_column + 1):
                raw_label = worksheet.cell(header_row, column).value
                label = str(raw_label).strip() if raw_label not in (None, "") else f"Coluna {get_column_letter(column)}"
                columns.append({"column": column, "label": label})

            rows = []
            for row_index in range(header_row + 1, end_row + 1):
                values = [
                    json_value(worksheet.cell(row_index, column).value)
                    for column in range(min_column, max_column + 1)
                ]
                if any(value not in (None, "") for value in values):
                    rows.append({"sourceRow": row_index, "values": values})

            relative_path = mapping_path.relative_to(root)
            discovered.append({
                "id": table_identifier(root, mapping_path, worksheet.title, header_row),
                "planName": plan_dir.name,
                "fileName": mapping_path.name,
                "relativePath": str(relative_path),
                "sheetName": worksheet.title,
                "headerRow": header_row,
                "columns": columns,
                "rows": rows,
                "rowCount": len(rows),
                "version": str(mapping_path.stat().st_mtime_ns),
                "_path": mapping_path,
                "_worksheet": worksheet,
                "_minColumn": min_column,
                "_maxColumn": max_column,
                "_endRow": end_row,
                "_excelTable": excel_table,
            })
    return discovered


def public_table(table):
    return {key: value for key, value in table.items() if not key.startswith("_")}


def list_tables(root_path):
    root = Path(root_path).resolve(strict=True)
    if not root.is_dir():
        raise ValueError("A pasta monitorada não foi encontrada.")

    tables = []
    errors = []
    for plan_dir in sorted((item for item in root.iterdir() if item.is_dir()), key=lambda item: item.name.lower()):
        mapping_path, _ = find_mapping_file(plan_dir)
        if mapping_path is None:
            continue
        workbook = None
        stream = None
        try:
            workbook, stream = open_workbook(mapping_path)
            found = discover_tables(workbook, root, plan_dir, mapping_path)
            if not found:
                errors.append({
                    "planName": plan_dir.name,
                    "fileName": mapping_path.name,
                    "message": "As colunas de código atual e novo código não foram encontradas.",
                })
            tables.extend(public_table(table) for table in found)
        except Exception as error:
            errors.append({
                "planName": plan_dir.name,
                "fileName": mapping_path.name,
                "message": str(error),
            })
        finally:
            if workbook is not None:
                workbook.close()
            if stream is not None:
                stream.close()

    return {"ok": True, "tables": tables, "errors": errors}


def parse_amount(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = str(value).strip().replace("\u00a0", " ")
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    filtered = "".join(character for character in text if character.isdigit() or character in ",.-")
    if not any(character.isdigit() for character in filtered):
        raise ValueError(f"Valor inválido: {value}")
    comma = filtered.rfind(",")
    dot = filtered.rfind(".")
    separator = None
    if comma >= 0 and dot >= 0:
        separator = "," if comma > dot else "."
    elif comma >= 0 and len(filtered) - comma - 1 in {1, 2}:
        separator = ","
    elif dot >= 0 and len(filtered) - dot - 1 in {1, 2}:
        separator = "."
    if separator:
        integer, decimal = filtered.rsplit(separator, 1)
        normalized = "".join(character for character in integer if character.isdigit() or character == "-")
        normalized = f"{normalized or '0'}.{''.join(character for character in decimal if character.isdigit())}"
    else:
        normalized = "".join(character for character in filtered if character.isdigit() or character == "-")
    if negative and not normalized.startswith("-"):
        normalized = f"-{normalized}"
    try:
        number = Decimal(normalized)
    except InvalidOperation as error:
        raise ValueError(f"Valor inválido: {value}") from error
    return int(number) if number == number.to_integral_value() else float(number)


def coerce_value(label, value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    if is_mapping_amount_header(label):
        return parse_amount(value)
    if is_current_header(label) or is_new_header(label):
        return str(value).strip()
    return value if isinstance(value, (int, float, bool)) else str(value).strip()


def atomic_save(workbook, mapping_path):
    suffix = ".xlsm" if mapping_path.suffix.lower() == ".xlsm" else ".xlsx"
    handle = tempfile.NamedTemporaryFile(prefix=f".{mapping_path.stem}-", suffix=suffix, dir=mapping_path.parent, delete=False)
    temporary_path = Path(handle.name)
    handle.close()
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, mapping_path)
    except PermissionError as error:
        raise PermissionError("Feche a planilha no Excel antes de salvar as alterações.") from error
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def save_table(root_path, payload):
    root = Path(root_path).resolve(strict=True)
    table_id = str(payload.get("tableId") or "")
    expected_version = str(payload.get("version") or "")
    submitted_rows = payload.get("rows")
    if not table_id or not isinstance(submitted_rows, list):
        raise ValueError("A tabela e suas linhas são obrigatórias.")
    if len(submitted_rows) > MAX_ROWS:
        raise ValueError("A tabela excede o limite de linhas permitido.")

    target_path = None
    target_plan = None
    for plan_dir in (item for item in root.iterdir() if item.is_dir()):
        mapping_path, _ = find_mapping_file(plan_dir)
        if mapping_path is None:
            continue
        workbook = None
        stream = None
        try:
            workbook, stream = open_workbook(mapping_path)
            matches = [table for table in discover_tables(workbook, root, plan_dir, mapping_path) if table["id"] == table_id]
            if matches:
                target_path = mapping_path
                target_plan = plan_dir
                break
        finally:
            if workbook is not None:
                workbook.close()
            if stream is not None:
                stream.close()
    if target_path is None:
        raise ValueError("A tabela TUSS selecionada não foi encontrada.")
    if not is_within(target_path.resolve(), root):
        raise ValueError("O arquivo TUSS está fora da pasta monitorada.")
    if expected_version and str(target_path.stat().st_mtime_ns) != expected_version:
        raise ValueError("A planilha foi alterada fora do painel. Atualize a tabela antes de salvar novamente.")

    workbook = None
    stream = None
    try:
        workbook, stream = open_workbook(target_path)
        matches = [table for table in discover_tables(workbook, root, target_plan, target_path) if table["id"] == table_id]
        if not matches:
            raise ValueError("A estrutura da tabela TUSS mudou. Atualize a página.")
        table = matches[0]
        columns = table["columns"]
        if not columns or len(columns) > MAX_COLUMNS:
            raise ValueError("A quantidade de colunas da tabela não é permitida.")
        rows = []
        for index, submitted in enumerate(submitted_rows, start=1):
            values = submitted.get("values") if isinstance(submitted, dict) else submitted
            if not isinstance(values, list) or len(values) != len(columns):
                raise ValueError(f"A linha {index} possui uma quantidade inválida de colunas.")
            rows.append([coerce_value(columns[column]["label"], value) for column, value in enumerate(values)])

        worksheet = table["_worksheet"]
        header_row = table["headerRow"]
        min_column = table["_minColumn"]
        max_column = table["_maxColumn"]
        new_end_row = header_row + len(rows)
        clear_end_row = max(table["_endRow"], new_end_row, header_row + 1)
        for row_index in range(header_row + 1, clear_end_row + 1):
            for column in range(min_column, max_column + 1):
                worksheet.cell(row_index, column).value = None
        for row_offset, values in enumerate(rows, start=1):
            for column_offset, value in enumerate(values):
                worksheet.cell(header_row + row_offset, min_column + column_offset).value = value

        excel_table = table["_excelTable"]
        if excel_table is not None:
            table_end_row = header_row + max(1, len(rows))
            excel_table.ref = (
                f"{get_column_letter(min_column)}{header_row}:"
                f"{get_column_letter(max_column)}{table_end_row}"
            )

        if stream is not None:
            stream.close()
            stream = None
        atomic_save(workbook, target_path)
    finally:
        if workbook is not None:
            workbook.close()
        if stream is not None:
            stream.close()

    refreshed = list_tables(root)
    saved_table = next((table for table in refreshed["tables"] if table["id"] == table_id), None)
    if saved_table is None:
        raise ValueError("A tabela foi salva, mas não pôde ser relida.")
    return {"ok": True, "table": saved_table}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["list", "save"])
    parser.add_argument("root_path")
    arguments = parser.parse_args()
    try:
        if arguments.mode == "list":
            result = list_tables(arguments.root_path)
        else:
            payload = json.load(sys.stdin)
            result = save_table(arguments.root_path, payload)
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
