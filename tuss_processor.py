import argparse
import hashlib
import json
import os
import re
import sys
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


WORKBOOK_EXTENSIONS = {".xlsx", ".xlsm"}
DEMONSTRATIVE_EXTENSIONS = WORKBOOK_EXTENSIONS | {".pdf"}
DEMONSTRATIVE_FOLDERS = {"demonstrativo xlsx", "demonstrativo pdf"}
OUTPUT_FOLDER = "TUSS Corrigidos"


def normalize_text(value):
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return re.sub(r"\s+", " ", text)


def normalize_code(value):
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip().upper()
    text = str(value).strip()
    if re.fullmatch(r"\d+[.,]0+", text):
        text = re.split(r"[.,]", text, maxsplit=1)[0]
    return re.sub(r"[^A-Z0-9]", "", unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii").upper())


def display_code(value):
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_current_header(value):
    key = normalize_text(value)
    return key in {"codigo atual", "cod atual", "tuss atual", "codigo tuss atual"} or "codigo atual" in key


def is_new_header(value):
    key = normalize_text(value)
    return key in {"novo codigo", "codigo novo", "novo tuss", "codigo tuss novo"} or "novo codigo" in key


def is_tuss_header(value):
    key = normalize_text(value)
    return key == "tuss" or key in {"codigo tuss", "cod tuss"}


def is_within(path, root):
    try:
        return os.path.commonpath([str(path), str(root)]) == str(root)
    except ValueError:
        return False


def workbook_options(path, read_only=False):
    return {
        "read_only": read_only,
        "data_only": False,
        "keep_vba": path.suffix.lower() == ".xlsm",
    }


def find_context(root_path, report_path):
    root = Path(root_path).resolve(strict=True)
    report = Path(report_path).resolve(strict=True)
    if not is_within(report, root):
        raise ValueError("O relatório selecionado está fora da pasta monitorada.")
    relative = report.relative_to(root)
    if len(relative.parts) < 3:
        raise ValueError("Não foi possível identificar o plano e a competência do relatório.")
    plan_dir = root / relative.parts[0]
    month_dir = plan_dir / relative.parts[1]
    if not month_dir.is_dir():
        raise ValueError("A pasta da competência selecionada não foi encontrada.")
    return root, report, plan_dir, month_dir


def find_mapping_file(plan_dir):
    candidates = []
    for item in plan_dir.iterdir():
        if not item.is_file():
            continue
        if "ajuste codigo tuss" not in normalize_text(item.stem):
            continue
        candidates.append(item)
    supported = [item for item in candidates if item.suffix.lower() in WORKBOOK_EXTENSIONS]
    exact = [item for item in supported if normalize_text(item.stem) == "ajuste codigo tuss"]
    if exact:
        return sorted(exact, key=lambda item: item.name.lower())[0], candidates
    return (sorted(supported, key=lambda item: item.name.lower())[0] if supported else None), candidates


def read_mapping(mapping_path):
    if openpyxl is None:
        raise RuntimeError("O módulo openpyxl não está instalado. Execute: pip install -r requirements.txt")
    workbook = openpyxl.load_workbook(mapping_path, **workbook_options(mapping_path, read_only=False))
    mappings = {}
    conflicts = []
    header_locations = []
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                current_columns = [cell.column for cell in row if is_current_header(cell.value)]
                new_columns = [cell.column for cell in row if is_new_header(cell.value)]
                if not current_columns or not new_columns:
                    continue
                current_column = current_columns[0]
                new_column = new_columns[0]
                header_locations.append({"sheet": worksheet.title, "row": row[0].row})
                for row_index in range(row[0].row + 1, worksheet.max_row + 1):
                    current_value = worksheet.cell(row_index, current_column).value
                    new_value = worksheet.cell(row_index, new_column).value
                    current_key = normalize_code(current_value)
                    new_key = normalize_code(new_value)
                    if not current_key or not new_key:
                        continue
                    entry = {
                        "current": display_code(current_value),
                        "currentKey": current_key,
                        "new": display_code(new_value),
                        "newKey": new_key,
                    }
                    existing = mappings.get(current_key)
                    if existing and existing["newKey"] != new_key:
                        conflicts.append({"current": entry["current"], "values": sorted({existing["new"], entry["new"]})})
                        continue
                    mappings[current_key] = entry
    finally:
        workbook.close()
    if not header_locations:
        raise ValueError("As colunas 'Código atual' e 'Novo Código' não foram encontradas no arquivo de ajuste.")
    return mappings, conflicts, header_locations


def find_demonstrative_files(month_dir):
    folders = []
    files = []
    output_key = normalize_text(OUTPUT_FOLDER)
    for candidate in month_dir.rglob("*"):
        if not candidate.is_dir() or normalize_text(candidate.name) not in DEMONSTRATIVE_FOLDERS:
            continue
        folders.append(candidate)
        for file_path in candidate.rglob("*"):
            if not file_path.is_file() or file_path.suffix.lower() not in DEMONSTRATIVE_EXTENSIONS:
                continue
            if output_key in {normalize_text(parent.name) for parent in file_path.parents}:
                continue
            if "tuss corrigido" in normalize_text(file_path.stem):
                continue
            files.append(file_path)
    unique_files = sorted(set(files), key=lambda item: str(item).lower())
    return sorted(set(folders), key=lambda item: str(item).lower()), unique_files


def summarize_matches(locations):
    counts = Counter((item["current"], item["new"]) for item in locations)
    return [
        {"current": current, "new": new, "count": count}
        for (current, new), count in sorted(counts.items())
    ]


def analyze_xlsx(file_path, mappings):
    if openpyxl is None:
        raise RuntimeError("O módulo openpyxl não está instalado. Execute: pip install -r requirements.txt")
    workbook = openpyxl.load_workbook(file_path, **workbook_options(file_path, read_only=False))
    locations = []
    columns = []
    seen_cells = set()
    try:
        for worksheet in workbook.worksheets:
            headers = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if is_tuss_header(cell.value):
                        headers.append((cell.row, cell.column))
                        columns.append({"sheet": worksheet.title, "column": cell.column, "headerRow": cell.row})
            for header_row, column in headers:
                for row_index in range(header_row + 1, worksheet.max_row + 1):
                    identity = (worksheet.title, row_index, column)
                    if identity in seen_cells:
                        continue
                    seen_cells.add(identity)
                    cell = worksheet.cell(row_index, column)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        continue
                    key = normalize_code(cell.value)
                    if key not in mappings:
                        continue
                    rule = mappings[key]
                    locations.append({
                        "sheet": worksheet.title,
                        "row": row_index,
                        "column": column,
                        "current": rule["current"],
                        "new": rule["new"],
                        "newKey": rule["newKey"],
                    })
    finally:
        workbook.close()
    return {
        "tussColumns": columns,
        "totalMatches": len(locations),
        "replacements": summarize_matches(locations),
        "_locations": locations,
        "status": "ready" if locations else ("no_matches" if columns else "no_tuss_column"),
    }


def pdf_words(document):
    all_pages = []
    header_centers = []
    for page_index, page in enumerate(document):
        words = page.get_text("words", sort=True)
        headers = []
        for word in words:
            if is_tuss_header(word[4]):
                center = (word[0] + word[2]) / 2
                headers.append(center)
                header_centers.append(center)
        all_pages.append((page_index, words, headers))
    return all_pages, header_centers


def analyze_pdf(file_path, mappings):
    try:
        import fitz
    except ImportError as error:
        raise RuntimeError("O módulo PyMuPDF não está instalado. Execute: pip install -r requirements.txt") from error
    document = fitz.open(file_path)
    locations = []
    try:
        if document.needs_pass:
            return {"tussColumns": [], "totalMatches": 0, "replacements": [], "_locations": [], "status": "protected"}
        pages, global_headers = pdf_words(document)
        for page_index, words, page_headers in pages:
            valid_headers = page_headers or global_headers
            for word in words:
                key = normalize_code(word[4])
                if key not in mappings or not valid_headers:
                    continue
                center = (word[0] + word[2]) / 2
                if min(abs(center - header) for header in valid_headers) > 80:
                    continue
                rule = mappings[key]
                locations.append({
                    "page": page_index,
                    "rect": [word[0], word[1], word[2], word[3]],
                    "current": rule["current"],
                    "new": rule["new"],
                    "newKey": rule["newKey"],
                })
    finally:
        document.close()
    columns = [{"header": "TUSS", "detected": True}] if global_headers else []
    return {
        "tussColumns": columns,
        "totalMatches": len(locations),
        "replacements": summarize_matches(locations),
        "_locations": locations,
        "status": "ready" if locations else ("no_matches" if columns else "no_tuss_column"),
    }


def analyze_file(file_path, mappings, month_dir):
    relative = file_path.relative_to(month_dir)
    base = {
        "id": hashlib.sha1(str(file_path).encode("utf-8")).hexdigest()[:12],
        "fileName": file_path.name,
        "relativePath": str(relative),
        "format": file_path.suffix.lower().lstrip(".").upper(),
        "_path": str(file_path),
    }
    try:
        details = analyze_pdf(file_path, mappings) if file_path.suffix.lower() == ".pdf" else analyze_xlsx(file_path, mappings)
        return {**base, **details}
    except Exception as error:
        return {**base, "tussColumns": [], "totalMatches": 0, "replacements": [], "_locations": [], "status": "error", "error": str(error)}


def build_analysis(root_path, report_path):
    root, _, plan_dir, month_dir = find_context(root_path, report_path)
    mapping_path, mapping_candidates = find_mapping_file(plan_dir)
    folders, demonstrative_files = find_demonstrative_files(month_dir)
    messages = []
    mappings = {}
    conflicts = []
    header_locations = []
    mapping_error = None
    if mapping_path:
        try:
            mappings, conflicts, header_locations = read_mapping(mapping_path)
        except Exception as error:
            mapping_error = str(error)
    elif mapping_candidates:
        mapping_error = "O arquivo de ajuste foi encontrado, mas o formato não é compatível. Use .xlsx ou .xlsm."
    else:
        messages.append("Arquivo 'Ajuste Codigo TUSS.xlsx' não encontrado na pasta do plano.")

    files = [analyze_file(file_path, mappings, month_dir) for file_path in demonstrative_files] if mappings and not conflicts else []
    if not folders:
        messages.append("Nenhuma pasta 'Demonstrativo XLSX' ou 'Demonstrativo PDF' foi encontrada nesta competência.")
    elif not demonstrative_files:
        messages.append("As pastas de demonstrativos não contêm arquivos XLSX, XLSM ou PDF.")
    if conflicts:
        messages.append("Existem códigos atuais repetidos com novos códigos diferentes no arquivo de ajuste.")
    if mapping_error:
        messages.append(mapping_error)

    total_matches = sum(item["totalMatches"] for item in files)
    ready = bool(mapping_path and mappings and not conflicts and demonstrative_files and not mapping_error)
    return {
        "ok": True,
        "planName": plan_dir.name,
        "periodName": month_dir.name,
        "mapping": {
            "found": bool(mapping_path),
            "fileName": mapping_path.name if mapping_path else None,
            "relativePath": str(mapping_path.relative_to(root)) if mapping_path else None,
            "totalRules": len(mappings),
            "rules": list(mappings.values()),
            "conflicts": conflicts,
            "headerLocations": header_locations,
            "_path": str(mapping_path) if mapping_path else None,
        },
        "demonstrativeFolders": [str(folder.relative_to(month_dir)) for folder in folders],
        "files": files,
        "totalFiles": len(demonstrative_files),
        "totalMatches": total_matches,
        "ready": ready,
        "messages": messages,
        "_root": str(root),
        "_monthDir": str(month_dir),
    }


def public_analysis(analysis):
    def clean(value):
        if isinstance(value, dict):
            return {key: clean(item) for key, item in value.items() if not key.startswith("_")}
        if isinstance(value, list):
            return [clean(item) for item in value]
        return value
    return clean(analysis)


def unique_output_path(output_dir, source_path):
    base = output_dir / f"{source_path.stem}_TUSS_CORRIGIDO{source_path.suffix.lower()}"
    if not base.exists():
        return base
    counter = 2
    while True:
        candidate = output_dir / f"{source_path.stem}_TUSS_CORRIGIDO_{counter}{source_path.suffix.lower()}"
        if not candidate.exists():
            return candidate
        counter += 1


def corrected_cell_value(original_value, new_code):
    if isinstance(original_value, int):
        return int(new_code) if new_code.isdigit() else new_code
    if isinstance(original_value, float) and original_value.is_integer():
        return int(new_code) if new_code.isdigit() else new_code
    return new_code


def apply_xlsx(source_path, target_path, locations):
    workbook = openpyxl.load_workbook(source_path, **workbook_options(source_path, read_only=False))
    try:
        for item in locations:
            cell = workbook[item["sheet"]].cell(item["row"], item["column"])
            cell.value = corrected_cell_value(cell.value, item["new"])
        temporary = target_path.with_name(f".{target_path.name}.tmp")
        workbook.save(temporary)
        os.replace(temporary, target_path)
    finally:
        workbook.close()


def apply_pdf(source_path, target_path, locations):
    import fitz
    document = fitz.open(source_path)
    try:
        grouped = {}
        for item in locations:
            grouped.setdefault(item["page"], []).append(item)
        for page_index, items in grouped.items():
            page = document[page_index]
            for item in items:
                rect = fitz.Rect(item["rect"])
                page.add_redact_annot(rect, fill=(1, 1, 1))
            page.apply_redactions(images=0, graphics=0, text=0)
            for item in items:
                rect = fitz.Rect(item["rect"])
                font_size = max(6, min(11, rect.height * 0.78))
                required_width = fitz.get_text_length(item["new"], fontname="helv", fontsize=font_size) + 3
                text_rect = fitz.Rect(rect.x0, rect.y0, max(rect.x1, rect.x0 + required_width), rect.y1 + 2)
                page.insert_textbox(text_rect, item["new"], fontname="helv", fontsize=font_size, color=(0, 0, 0), align=0)
        temporary = target_path.with_name(f".{target_path.name}.tmp")
        document.save(temporary, garbage=4, deflate=True)
        os.replace(temporary, target_path)
    finally:
        document.close()


def apply_corrections(root_path, report_path):
    analysis = build_analysis(root_path, report_path)
    if not analysis["ready"]:
        raise ValueError("Os arquivos necessários para o ajuste TUSS não estão disponíveis ou possuem erros.")
    if analysis["totalMatches"] == 0:
        return {"ok": True, "applied": False, "totalReplacements": 0, "outputs": [], "message": "Nenhum código precisa ser ajustado."}

    month_dir = Path(analysis["_monthDir"])
    output_dir = month_dir / OUTPUT_FOLDER
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = []
    for file_info in analysis["files"]:
        if not file_info["totalMatches"]:
            continue
        source_path = Path(file_info["_path"])
        target_path = unique_output_path(output_dir, source_path)
        if source_path.suffix.lower() == ".pdf":
            apply_pdf(source_path, target_path, file_info["_locations"])
        else:
            apply_xlsx(source_path, target_path, file_info["_locations"])
        outputs.append({
            "source": file_info["relativePath"],
            "output": str(target_path.relative_to(month_dir)),
            "format": file_info["format"],
            "replacements": file_info["totalMatches"],
            "changes": file_info["replacements"],
        })

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    audit_path = output_dir / f"registro_ajustes_TUSS_{timestamp}.json"
    audit = {
        "createdAt": datetime.now().isoformat(timespec="seconds"),
        "plan": analysis["planName"],
        "period": analysis["periodName"],
        "mappingFile": analysis["mapping"]["fileName"],
        "totalReplacements": sum(item["replacements"] for item in outputs),
        "outputs": outputs,
    }
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "ok": True,
        "applied": True,
        "totalReplacements": audit["totalReplacements"],
        "outputs": outputs,
        "auditFile": str(audit_path.relative_to(month_dir)),
        "outputFolder": OUTPUT_FOLDER,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["analyze", "apply"])
    parser.add_argument("root_path")
    parser.add_argument("report_path")
    arguments = parser.parse_args()
    try:
        if arguments.mode == "analyze":
            result = public_analysis(build_analysis(arguments.root_path, arguments.report_path))
        else:
            result = apply_corrections(arguments.root_path, arguments.report_path)
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
