import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

import fitz
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tuss_processor import apply_corrections, build_analysis


class TussProcessorTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name) / "DemonstPagProjeto"
        self.plan = self.root / "Plano Teste"
        self.month = self.plan / "AGOSTO_2026"
        self.demo_folder = self.month / "Demonstrativo XLSX"
        self.demo_folder.mkdir(parents=True)
        self.report = self.month / "relatorio_teste.txt"
        self.report.write_text("RELATORIO DE DEMONSTRATIVOS DE PAGAMENTO", encoding="utf-8")

    def tearDown(self):
        self.temporary.cleanup()

    def create_mapping(self):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "TUSS"
        worksheet.append(["Código atual", "Novo Código"])
        worksheet.append(["XXX", 64617270])
        workbook.save(self.plan / "Ajuste Codigo TUSS.xlsx")
        workbook.close()

    def test_xlsx_replaces_every_tuss_column_and_preserves_original(self):
        self.create_mapping()
        source = self.demo_folder / "demonstrativo.xlsx"
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Dados"
        worksheet.append(["Guia", "TUSS", "Valor", "TUSS"])
        worksheet.append(["A", "XXX", 10, "XXX"])
        worksheet.append(["B", "64617270", 20, "XXX"])
        workbook.save(source)
        workbook.close()
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

        analysis = build_analysis(self.root, self.report)
        self.assertTrue(analysis["ready"])
        self.assertEqual(analysis["totalMatches"], 3)

        result = apply_corrections(self.root, self.report)
        self.assertTrue(result["applied"])
        self.assertEqual(result["totalReplacements"], 3)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)

        corrected = self.month / result["outputs"][0]["output"]
        corrected_book = openpyxl.load_workbook(corrected, read_only=False)
        corrected_sheet = corrected_book["Dados"]
        self.assertEqual(corrected_sheet["B2"].value, "64617270")
        self.assertEqual(corrected_sheet["D2"].value, "64617270")
        self.assertEqual(corrected_sheet["D3"].value, "64617270")
        corrected_book.close()
        audit = self.month / result["auditFile"]
        self.assertEqual(json.loads(audit.read_text(encoding="utf-8"))["totalReplacements"], 3)

    def test_missing_mapping_is_reported_without_writing(self):
        analysis = build_analysis(self.root, self.report)
        self.assertFalse(analysis["ready"])
        self.assertFalse(analysis["mapping"]["found"])
        self.assertFalse((self.month / "TUSS Corrigidos").exists())

    def test_pdf_replaces_only_values_aligned_with_tuss_column(self):
        self.create_mapping()
        pdf_folder = self.month / "Demonstrativo PDF"
        pdf_folder.mkdir()
        source = pdf_folder / "demonstrativo.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_text((80, 80), "TUSS", fontsize=10)
        page.insert_text((80, 105), "XXX", fontsize=10)
        page.insert_text((260, 105), "XXX", fontsize=10)
        document.save(source)
        document.close()
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

        analysis = build_analysis(self.root, self.report)
        pdf_analysis = next(item for item in analysis["files"] if item["format"] == "PDF")
        self.assertEqual(pdf_analysis["totalMatches"], 1)

        result = apply_corrections(self.root, self.report)
        pdf_output = next(item for item in result["outputs"] if item["format"] == "PDF")
        corrected = self.month / pdf_output["output"]
        corrected_document = fitz.open(corrected)
        text = "\n".join(page.get_text() for page in corrected_document)
        corrected_document.close()
        self.assertIn("64617270", text)
        self.assertIn("XXX", text)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)


if __name__ == "__main__":
    unittest.main()
