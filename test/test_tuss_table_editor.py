import tempfile
import unittest
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tuss_table_editor import list_tables, save_table


class TussTableEditorTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name) / "DemonstPagProjeto"
        self.plan = self.root / "Sulamerica"
        self.plan.mkdir(parents=True)
        (self.root / "Intermedica").mkdir()
        self.mapping = self.plan / "Ajuste Codigo TUSS.xlsx"

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "TUSS"
        worksheet.append(["Codigo Atual", "Novo Codigo", "Valor", "Descrição"])
        worksheet.append(["XXX", 64500039, 4600, "Pacote A"])
        worksheet.append(["YYY", 64617270, 276.2, "Pacote B"])
        worksheet["A1"].fill = PatternFill("solid", fgColor="0D766F")
        table = Table(displayName="Tabela1", ref="A1:D3")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        worksheet.add_table(table)
        workbook.save(self.mapping)
        workbook.close()

    def tearDown(self):
        self.temporary.cleanup()

    def test_lists_only_plans_with_a_valid_mapping_table(self):
        result = list_tables(self.root)

        self.assertTrue(result["ok"])
        self.assertEqual(result["errors"], [])
        self.assertEqual(len(result["tables"]), 1)
        table = result["tables"][0]
        self.assertEqual(table["planName"], "Sulamerica")
        self.assertEqual([column["label"] for column in table["columns"]], [
            "Codigo Atual", "Novo Codigo", "Valor", "Descrição"
        ])
        self.assertEqual(table["rowCount"], 2)

    def test_edits_and_deletes_rows_while_preserving_the_excel_table(self):
        table = list_tables(self.root)["tables"][0]
        result = save_table(self.root, {
            "tableId": table["id"],
            "version": table["version"],
            "rows": [{"values": ["ABC", "99999999", "123,45", "Descrição atualizada"]}],
        })

        self.assertTrue(result["ok"])
        self.assertEqual(result["table"]["rowCount"], 1)
        workbook = openpyxl.load_workbook(self.mapping)
        worksheet = workbook["TUSS"]
        self.assertEqual(worksheet["A2"].value, "ABC")
        self.assertEqual(worksheet["B2"].value, "99999999")
        self.assertEqual(worksheet["C2"].value, 123.45)
        self.assertEqual(worksheet["D2"].value, "Descrição atualizada")
        self.assertIsNone(worksheet["A3"].value)
        self.assertEqual(worksheet.tables["Tabela1"].ref, "A1:D2")
        self.assertEqual(worksheet["A1"].fill.fgColor.rgb[-6:], "0D766F")
        workbook.close()

    def test_accepts_the_established_xslx_file_name(self):
        renamed = self.mapping.with_suffix(".xslx")
        self.mapping.rename(renamed)

        result = list_tables(self.root)

        self.assertEqual(len(result["tables"]), 1)
        self.assertEqual(result["tables"][0]["fileName"], "Ajuste Codigo TUSS.xslx")

    def test_rejects_a_stale_file_version(self):
        table = list_tables(self.root)["tables"][0]
        with self.assertRaisesRegex(ValueError, "alterada fora do painel"):
            save_table(self.root, {
                "tableId": table["id"],
                "version": "1",
                "rows": table["rows"],
            })


if __name__ == "__main__":
    unittest.main()
